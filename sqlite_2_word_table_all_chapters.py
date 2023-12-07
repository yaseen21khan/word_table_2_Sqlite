import docx
import sqlite3

def export_sqlite_to_word(db_file, output_word_file):
    conn = sqlite3.connect(db_file)
    cursor = conn.cursor()

    # Fetch data from the SQLite table
    cursor.execute('SELECT * FROM YourTable')
    data = cursor.fetchall()

    conn.close()

    # Create a Word document
    doc = docx.Document()

    # Add a table to the Word document
    table = doc.add_table(rows=1, cols=4)

    # Add headers to the table
    headers = ["Level", "Chapter1", "Korean", "English"]
    for col, header in enumerate(headers):
        table.cell(0, col).text = header

    # Add data to the table
    for row in data:
        new_row = table.add_row()
        for col, value in enumerate(row):
            new_row.cells[col].text = str(value)
            print(value)

    # Save the Word document
    doc.save(output_word_file)

    print(f"Data has been successfully exported to {output_word_file}.")

if __name__ == "__main__":
    sqlite_db_file = "D:/Level2_vocabulary_database.db"
    output_word_file = "D:/Level2_vocabulary.docx"

    # Export data from SQLite to Word
    export_sqlite_to_word(sqlite_db_file, output_word_file)
